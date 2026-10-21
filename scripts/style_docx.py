# Designed by TiantianYe
# -*- coding: utf-8 -*-
"""pandoc docx 样式后处理: 字体/字号/标题色/页边距/表格边框与表头底纹/图片限宽/条目间隔"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
doc = Document(SRC)

def set_style_font(style, ascii_font, ea_font, size, bold=None, color=None):
    style.font.name = ascii_font
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)

S = doc.styles
# 正文
try:
    normal = S["Normal"]
    set_style_font(normal, "微软雅黑", "微软雅黑", 10.5)
    normal.paragraph_format.line_spacing = 1.6
    normal.paragraph_format.space_after = Pt(4)
except KeyError:
    pass
# 标题
for name, size, color, before in [("Heading 1", 17, "1A237E", 0), ("Heading 2", 13, "1A237E", 0), ("Heading 3", 11.5, "212121", 10), ("Heading 4", 11, "212121", 10), ("Title", 22, "1A237E", 0)]:
    try:
        st = S[name]
        set_style_font(st, "微软雅黑", "微软雅黑", size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(8 if name in ("Heading 1", "Heading 2", "Title") else 4)
    except KeyError:
        pass

# 页边距
for sec in doc.sections:
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

# 页眉：Oli教室专用（每一页）
from docx.enum.text import WD_ALIGN_PARAGRAPH as _WAL
for sec in doc.sections:
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    for r in list(hp.runs):
        r._element.getparent().remove(r._element)
    run = hp.add_run("Oli教室专用")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("8B97A8")
    run.font.name = "微软雅黑"
    _rpr = run._element.get_or_add_rPr()
    _rf = _rpr.find(qn("w:rFonts"))
    if _rf is None:
        _rf = OxmlElement("w:rFonts")
        _rpr.append(_rf)
    _rf.set(qn("w:ascii"), "微软雅黑")
    _rf.set(qn("w:hAnsi"), "微软雅黑")
    _rf.set(qn("w:eastAsia"), "微软雅黑")
    hp.alignment = _WAL.CENTER

# 引用块（原题块）样式：左缩进区分，不加底色
try:
    bt = S["Block Text"]
    set_style_font(bt, "微软雅黑", "微软雅黑", 10.5)
    bt.paragraph_format.left_indent = Cm(0.5)
    bt.paragraph_format.right_indent = Cm(0.3)
    bt.paragraph_format.space_before = Pt(4)
    bt.paragraph_format.space_after = Pt(4)
    bt.paragraph_format.line_spacing = 1.5
except KeyError:
    pass

# 表格：全边框 + 表头底纹加粗 + 表内字号 9.5pt
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "E0E0E0")
        borders.append(el)
    tblPr.append(borders)

for table in doc.tables:
    set_table_borders(table)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            if ri == 0:
                set_cell_bg(cell, "F5F7FA")
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.4
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string("1A237E")

# ---------- 排版守护 ----------
# 1) 标题与下一段同页、标题自身不跨页（防止标题孤悬页底）
for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"):
    try:
        pf = S[name].paragraph_format
        pf.keep_with_next = True
        pf.keep_together = True
    except KeyError:
        pass

# 2) 图片段落【不要】设置 keep_with_next（用户 2026-10-21 反馈）：
#    keepNext 会让 Word 在段落左侧显示黑色小方点（与标题相同），图片看起来像"以标题格式插入"；
#    成品报告（旧脚本）图片无此标记。图片与答案跨页分离可接受，禁止恢复此设置。

# 3) 行内公式后紧跟的标点前插入词连接符 U+2060（防止 。，出现在行首）
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WJ = "\u2060"
PUNCT = set("。，、；：）")
n_wj = 0
for om in list(doc.element.body.iter(f"{{{M_NS}}}oMath")):
    parent = om.getparent()
    if parent is None:
        continue
    sibs = list(parent)
    i = sibs.index(om)
    if i + 1 < len(sibs):
        nxt = sibs[i + 1]
        if nxt.tag == f"{{{W_NS}}}r":
            for t in nxt.findall(f"{{{W_NS}}}t"):
                if t.text and t.text[0] in PUNCT:
                    t.text = WJ + t.text
                    n_wj += 1
print("WJ 插入标点数:", n_wj)

# ---------- 章节间隔：真实回车空段落（大章节前 2 行、小章节前 1 行、变式题间 1 行） ----------
_blanks_made = {"h1": 0, "h2": 0, "h4": 0}
_paras = list(doc.paragraphs)
_first_seen = False
_first_h4_seen = False
for _para in _paras:
    _sname = _para.style.name if _para.style is not None else ""
    if _sname in ("Heading 1", "Heading 2"):
        if not _first_seen:
            _first_seen = True   # 文档第一个标题前不加空行
            continue
        _n = 2 if _sname == "Heading 1" else 1
        for _ in range(_n):
            _blank = OxmlElement("w:p")
            _para._p.addprevious(_blank)
            _blanks_made["h1" if _sname == "Heading 1" else "h2"] += 1
    elif _sname == "Heading 4":
        # 第七部分变式题：相邻两道题之间空一行（Heading 4 = 变式标题，首道跳过）
        if not _first_h4_seen:
            _first_h4_seen = True
            continue
        _blank = OxmlElement("w:p")
        _para._p.addprevious(_blank)
        _blanks_made["h4"] += 1
print("空段落: H1 前", _blanks_made["h1"], "| H2 前", _blanks_made["h2"], "| H4(变式) 前", _blanks_made["h4"])

# ---------- 题目条目间隔：每个条目头段前 1 个空段落（小节首条跳过） ----------
# 条目头判定（与内容源模板对应，全部为"加粗起头"段）：
#   **（N）考点名…**            （3.2/3.3/5.4 题图集等编号条目）
#   **典型原题（卷名 第 X 题）…** （2.x 选择题条目）
#   **…——代表题：卷名 第 X 题**  （4.2 高频题型拆解等）
# 引用块（Block Text）内的加粗题干行、各级标题、答案行均不算条目头。
_HEAD_RE = re.compile(r"^（\d+）|^典型原题")

def _is_entry_head(p):
    if p.style is None or p.style.name in ("Block Text", "Quote"):
        return False
    txt = p.text.strip()
    if not txt:
        return False
    first = next((r for r in p.runs if r.text.strip()), None)
    if first is None or not first.bold:
        return False
    if _HEAD_RE.match(txt):
        return True
    if "——代表题：" in txt or "——代表题:" in txt:
        return True
    return False

_n_entry = 0
_new_section = True   # 文档开头与每个标题之后视为新小节，首条跳过
for _para in list(doc.paragraphs):
    _sname = _para.style.name if _para.style is not None else ""
    if _sname.startswith("Heading") or _sname == "Title":
        _new_section = True
        continue
    if _is_entry_head(_para):
        if not _new_section:
            _blank = OxmlElement("w:p")
            _para._p.addprevious(_blank)
            _n_entry += 1
        _new_section = False
print("条目间隔空段（题目间换行）:", _n_entry)

# ---------- 图片尺寸守护：超限等比缩小（默认宽/高上限 6cm） ----------
# pandoc 按像素 96dpi 原生尺寸插图，题图常被放大到 9cm 级（原卷自然宽度多在 2.5-4.5cm）。
# 可用环境变量 IMG_MAX_W_CM / IMG_MAX_H_CM 覆盖上限。
MAX_W = Cm(float(os.environ.get("IMG_MAX_W_CM", "6")))
MAX_H = Cm(float(os.environ.get("IMG_MAX_H_CM", "6")))
_n_shrunk = 0
for _shape in doc.inline_shapes:
    try:
        _w, _h = _shape.width, _shape.height
    except Exception:
        continue
    if not _w or not _h:
        continue
    _scale = min(MAX_W / _w, MAX_H / _h, 1.0)
    if _scale < 1.0:
        _shape.width = Emu(int(_w * _scale))
        _shape.height = Emu(int(_h * _scale))
        _n_shrunk += 1
print("图片缩小:", _n_shrunk, "张 (上限 %.1fx%.1fcm)" % (MAX_W.cm, MAX_H.cm))

doc.save(SRC)
print("styled:", SRC)
print("tables:", len(doc.tables), "| paragraphs:", len(doc.paragraphs))
